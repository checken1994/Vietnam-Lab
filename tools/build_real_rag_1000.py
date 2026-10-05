import json, re, time, hashlib
import logging
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
import requests
import openpyxl

logger = logging.getLogger(__name__)

ROOT=Path(__file__).resolve().parents[1]
SRC=ROOT/'benchmark'/'bo_de_1000_cau_v3-v2.xlsx'
OUT_JSONL=ROOT/'benchmark'/'questions_1000_real_rag_20260817.jsonl'
OUT_XLSX=ROOT/'benchmark'/'bo_de_1000_cau_real_rag_20260817.xlsx'
UA='SCP-Real-RAG-Benchmark/1.0 (local benchmark)'

def clean_query(q):
    q=re.sub(r'\(Mã\s*CH-\d+\)', '', q, flags=re.I)
    q=re.sub(r'\s+', ' ', q.replace('\n',' ')).strip(' ?')
    q=re.sub(r'\b(mới nhất|hiện nay|năm nay)\b','',q,flags=re.I)
    return re.sub(r'\s+',' ',q).strip(' ?')

def api(session, lang, params):
    url=f'https://{lang}.wikipedia.org/w/api.php'
    for attempt in range(3):
        try:
            r=session.get(url,params=params,headers={'User-Agent':UA},timeout=20)
            r.raise_for_status(); return r.json()
        except Exception as exc:
            logger.debug("wikipedia api call failed (attempt %d)", attempt + 1, exc_info=exc)
            if attempt==2: return {}
            time.sleep(1.2*(attempt+1))

def fetch_one(item):
    idx,q=item; s=requests.Session(); query=clean_query(q); candidates=[]
    for lang in ('vi','en'):
        data=api(s,lang,{'action':'query','list':'search','srsearch':query,'srlimit':5,'format':'json','utf8':1})
        for x in data.get('query',{}).get('search',[]):
            title=x.get('title','')
            if title and title not in [c[1] for c in candidates]: candidates.append((lang,title))
        if candidates: break
    docs=[]
    for lang,title in candidates[:5]:
        data=api(s,lang,{'action':'query','prop':'extracts|info','explaintext':1,'exintro':1,'inprop':'url','titles':title,'format':'json','utf8':1})
        page=next(iter(data.get('query',{}).get('pages',{}).values()),{})
        text=(page.get('extract') or '').strip(); url=page.get('fullurl') or f'https://{lang}.wikipedia.org/wiki/{title.replace(" ","_")}'
        if text: docs.append({'lang':lang,'title':title,'url':url,'text':text[:10000]})
    base={'index':idx,'question_id':f'CH-{idx:04d}','question':q,'answer':'','domain':'general','source_date':'retrieved_2026-08-17'}
    if not docs:
        base.update({'contexts':[],'ground_truth':'','source_url':'','ground_truth_source_url':'','retrieval_status':'NO_DOCUMENT','review_status':'manual_review_required'}); return base
    primary=docs[0]; secondary=docs[1] if len(docs)>1 else docs[0]
    chunk=hashlib.sha256((primary['url']+primary['text']).encode()).hexdigest()[:16]
    ctx=f"[chunk_id=wiki-{chunk}] source_title={primary['title']} source_url={primary['url']} language={primary['lang']}\n{primary['text']}"
    gt=f"[ground_truth_source={secondary['url']}] {secondary['text'][:6000]}"
    base.update({'contexts':[ctx],'ground_truth':gt,'source_url':primary['url'],'ground_truth_source_url':secondary['url'],'source_title':primary['title'],'ground_truth_title':secondary['title'],'retrieval_status':'OK','review_status':'auto_two_document_review'}); return base

def main():
    wb=openpyxl.load_workbook(SRC,read_only=True,data_only=True); rows=list(wb.active.iter_rows(values_only=True))[1:]
    items=[(i+1,str(r[0] or '').strip()) for i,r in enumerate(rows) if str(r[0] or '').strip()]; results=[None]*len(items)
    with ThreadPoolExecutor(max_workers=8) as ex:
        futs={ex.submit(fetch_one,item):item[0] for item in items}; done=0
        for f in as_completed(futs):
            results[futs[f]-1]=f.result(); done+=1
            if done%50==0: print('retrieved',done,flush=True)
    OUT_JSONL.write_text('\n'.join(json.dumps(x,ensure_ascii=False) for x in results)+'\n',encoding='utf-8')
    out=openpyxl.Workbook(); ws=out.active; ws.title='1000_Cau_Hoi_Real_RAG'; cols=['question_id','question','contexts','answer','ground_truth','source_url','ground_truth_source_url','domain','retrieval_status','review_status','source_date']; ws.append(cols)
    for x in results: ws.append([json.dumps(x.get(k,[]),ensure_ascii=False) if k=='contexts' else x.get(k,'') for k in cols])
    out.save(OUT_XLSX); ok=sum(x['retrieval_status']=='OK' for x in results); print(json.dumps({'rows':len(results),'retrieval_ok':ok,'no_document':len(results)-ok,'jsonl':str(OUT_JSONL),'xlsx':str(OUT_XLSX)},ensure_ascii=False))
if __name__=='__main__': main()
